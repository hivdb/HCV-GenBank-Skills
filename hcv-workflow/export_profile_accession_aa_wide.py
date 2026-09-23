#!/usr/bin/env python3
"""Export QC-passed accession amino-acid calls in wide format."""

from __future__ import annotations

import argparse
import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook


CANONICAL_AAS = frozenset("ACDEFGHIKLMNPQRSTVWY*")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-workbook", type=Path, required=True)
    parser.add_argument("--profile-accessions-csv", type=Path, required=True)
    parser.add_argument("--start-position", type=int, required=True)
    parser.add_argument("--end-position", type=int, required=True)
    parser.add_argument("--output-csv", type=Path, required=True)
    parser.add_argument(
        "--copy-output-csv",
        type=Path,
        help="Optional repository-level copy of the completed CSV.",
    )
    return parser.parse_args()


def load_profile_accessions(path: Path) -> set[str]:
    with path.open(encoding="utf-8", newline="") as handle:
        return {
            str(row.get("accession") or "").strip()
            for row in csv.DictReader(handle)
            if str(row.get("accession") or "").strip()
        }


def normalize_genotype(value: object) -> str:
    genotype = str(value or "").strip()
    return genotype if genotype.upper().startswith("GT") else f"GT{genotype}"


def export_wide_table(args: argparse.Namespace) -> dict[str, int]:
    if args.end_position < args.start_position:
        raise ValueError("--end-position must be greater than or equal to --start-position")
    positions = range(args.start_position, args.end_position + 1)
    allowed_accessions = load_profile_accessions(args.profile_accessions_csv)
    workbook = load_workbook(args.input_workbook, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    header = [str(value or "") for value in next(worksheet.iter_rows(values_only=True))]
    index = {name: position for position, name in enumerate(header)}
    required = [
        "AccessionID",
        "ClosestGT",
        "ClosestSubtype",
        "StartAAPosition",
        "AASequence",
    ]
    missing = [name for name in required if name not in index]
    if missing:
        raise RuntimeError(f"Missing columns in profile input: {', '.join(missing)}")

    rows: dict[str, dict[str, str]] = {}
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        accession = str(values[index["AccessionID"]] or "").strip()
        if accession not in allowed_accessions:
            continue
        if (
            "AlignmentQCStatus" in index
            and str(values[index["AlignmentQCStatus"]] or "").strip() != "PASS"
        ):
            continue
        start = values[index["StartAAPosition"]]
        sequence = str(values[index["AASequence"]] or "").strip().upper()
        if start in (None, "") or not sequence:
            continue
        row = rows.setdefault(
            accession,
            {
                "Accession": accession,
                "Genotype": normalize_genotype(values[index["ClosestGT"]]),
                "Subtype": str(values[index["ClosestSubtype"]] or "").strip(),
            },
        )
        for offset, amino_acid in enumerate(sequence):
            position = int(start) + offset
            if position in positions and amino_acid in CANONICAL_AAS:
                key = str(position)
                previous = row.get(key)
                if previous and previous != amino_acid:
                    raise RuntimeError(
                        f"Conflicting amino-acid calls for {accession} at position {position}"
                    )
                row[key] = amino_acid
    workbook.close()

    fieldnames = ["Accession", "Genotype", "Subtype"] + [str(position) for position in positions]
    args.output_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.output_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows.values():
            writer.writerow(row)
    if args.copy_output_csv:
        args.copy_output_csv.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(args.output_csv, args.copy_output_csv)
    return {"accession_count": len(rows), "position_count": len(positions)}


def main() -> int:
    print(export_wide_table(parse_args()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
