#!/usr/bin/env python3
"""Merge NS3 subtype complete-profile worksheets into one simplified table."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook


COLUMNS = [
    "Subtype",
    "NS3Position",
    "NumSeqsIncludingPosition",
    "AminoAcid",
    "CountWithAA",
    "PctWithAA",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-workbook", required=True)
    parser.add_argument("--output-workbook", required=True)
    parser.add_argument("--range-output-workbook")
    parser.add_argument("--range-start", type=int)
    parser.add_argument("--range-end", type=int)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source = load_workbook(args.input_workbook, read_only=True, data_only=True)
    range_output = Workbook() if args.range_output_workbook else None
    range_sheet = range_output.active if range_output else None
    if range_sheet:
        range_sheet.title = "NS3_Subtype_CompleteProfiles"
        range_sheet.append(COLUMNS)
    output = Workbook()
    sheet = output.active
    sheet.title = "NS3_Subtype_CompleteProfiles"
    sheet.append(COLUMNS)
    source_sheet_count = len(source.sheetnames)
    merged_rows = 0
    for source_sheet in source.worksheets:
        header = [
            str(value or "") for value in next(source_sheet.iter_rows(values_only=True))
        ]
        index = {name: position for position, name in enumerate(header)}
        missing = [name for name in COLUMNS if name not in index]
        if missing:
            raise RuntimeError(
                f"Columns missing from worksheet {source_sheet.title}: {', '.join(missing)}"
            )
        for values in source_sheet.iter_rows(min_row=2, values_only=True):
            row = [values[index[name]] for name in COLUMNS]
            sheet.append(row)
            if range_sheet is not None:
                position = row[1]
                if ((args.range_start is None or position >= args.range_start)
                        and (args.range_end is None or position <= args.range_end)):
                    range_sheet.append(row)
            merged_rows += 1
    source.close()
    output_path = Path(args.output_workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output.save(output_path)
    range_rows = None
    if range_output is not None:
        range_path = Path(args.range_output_workbook)
        range_path.parent.mkdir(parents=True, exist_ok=True)
        range_output.save(range_path)
        range_rows = range_sheet.max_row - 1
    print(
        json.dumps(
            {
                "output_workbook": str(output_path.resolve()),
                "source_sheet_count": source_sheet_count,
                "merged_row_count": merged_rows,
                "range_output_workbook": str(Path(args.range_output_workbook).resolve()) if args.range_output_workbook else None,
                "range_row_count": range_rows,
            }
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
