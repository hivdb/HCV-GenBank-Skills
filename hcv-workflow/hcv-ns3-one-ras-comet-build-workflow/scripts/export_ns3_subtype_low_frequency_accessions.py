#!/usr/bin/env python3
"""Link observed subtype RAS calls below 5% to source accessions and RefIDs."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from build_ns3_completeprofiles_tabspergt import load_rows

RAS_POSITIONS = {36, 41, 43, 54, 55, 56, 80, 122, 155, 156, 158, 166, 168, 170, 175}
REFERENCE_COLUMNS = ["RefName", "RefYear", "PMID", "Title", "Author", "Journal"]
SPECIAL_TARGETS = [("1a", 155, "K"), ("3b", 155, "MI"), ("6l", 156, "V")]
COLUMNS = ["Genotype", "Subtype", "NS3Position", "AminoAcid", "CountWithAA",
           "NumSeqsIncludingPosition", "PctWithAA", "AccessionID", "RefID", *REFERENCE_COLUMNS]


def load_references(path: Path, sheet_name: str) -> dict[str, tuple[str, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        records = workbook[sheet_name].iter_rows(values_only=True)
        header = next(records)
        fields = ["RefName", "RefYear" if "RefYear" in header else "Year", "PMID", "Title", "Author", "Journal"]
        missing = set(["RefID", *fields]) - set(header)
        if missing:
            raise ValueError(f"Missing reference columns: {sorted(missing)}")
        references = {}
        for values in records:
            row = dict(zip(header, values))
            ref = str(row["RefID"] or "").strip()
            if not ref:
                continue
            metadata = tuple(str(row[field]) if row[field] is not None else "" for field in fields)
            if ref in references and references[ref] != metadata:
                raise ValueError(f"Conflicting reference metadata for RefID {ref}")
            references[ref] = metadata
        return references
    finally:
        workbook.close()


def export(profile_path: Path, source_path: Path, output_path: Path,
           reference_path: Path, reference_sheet: str) -> int:
    references = load_references(reference_path, reference_sheet)
    rare = {}
    workbook = load_workbook(profile_path, read_only=True, data_only=True)
    for sheet in workbook:
        genotype = sheet.title.removeprefix("GT")
        records = sheet.iter_rows(values_only=True)
        header = next(records)
        for values in records:
            row = dict(zip(header, values))
            position = int(row["NS3Position"])
            aa = str(row["AminoAcid"])
            count, total = int(row["CountWithAA"]), int(row["NumSeqsIncludingPosition"])
            if position in RAS_POSITIONS and aa not in {"X", "*"} and 0 < 100 * count < 5 * total:
                rare[(genotype, str(row["Subtype"]), position, aa)] = (count, total)
    workbook.close()

    # Reuse profile construction's QC, assignment, and callable-RAS gates.
    eligible, _ = load_rows(source_path)
    refs = defaultdict(set)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    records = workbook.worksheets[0].iter_rows(values_only=True)
    header = next(records)
    if "RefID" not in header:
        raise ValueError("Profile source workbook is missing RefID")
    for values in records:
        row = dict(zip(header, values))
        refs[str(row["AccessionID"]).strip()].add(str(row["RefID"] or "").strip())
    workbook.close()

    output = set()
    matched = set()
    for row in eligible:
        for position in RAS_POSITIONS:
            offset = position - row["StartAAPosition"]
            if not 0 <= offset < len(row["AASequence"]):
                continue
            aa = row["AASequence"][offset]
            key = (row["ClosestGT"], row["ClosestSubtype"], position, aa)
            if key not in rare:
                continue
            matched.add(key)
            count, total = rare[key]
            for ref in refs[row["AccessionID"]]:
                output.add((*key, count, total, 100.0 * count / total, row["AccessionID"], ref))
    if rare.keys() - matched:
        raise ValueError(f"No eligible source accessions for {len(rare.keys() - matched)} rare profile calls")
    missing_refs = {row[-1] for row in output} - references.keys()
    if missing_refs:
        raise ValueError(f"RefIDs missing from original reference worksheet: {sorted(missing_refs)}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(COLUMNS)
        writer.writerows((*row, *references[row[-1]]) for row in sorted(output))
    for subtype, position, amino_acids in SPECIAL_TARGETS:
        special_rows = [row for row in sorted(output)
                        if (row[1], row[2]) == (subtype, position) and row[3] in amino_acids]
        special_path = output_path.with_name(
            f"NS3_Subtype_RAS_Below5Pct_{subtype}_{position}{'_or_'.join(amino_acids)}_Accessions_Refs.xlsx"
        )
        write_special_workbook(special_path, special_rows, references)
    return len(output)


def write_special_workbook(path: Path, rows: list[tuple], references: dict) -> None:
    workbook = Workbook()
    accessions = workbook.active
    accessions.title = "Accessions"
    accessions.append(COLUMNS[:9])
    by_reference = defaultdict(set)
    for row in rows:
        accessions.append(row)
        by_reference[row[-1]].add(row[-2])
    reference_sheet = workbook.create_sheet("References")
    reference_sheet.append(["RefID", *REFERENCE_COLUMNS, "AccessionCount"])
    for ref, accession_ids in sorted(by_reference.items()):
        reference_sheet.append([ref, *references[ref], len(accession_ids)])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            sheet.column_dimensions[get_column_letter(cell.column)].width = (
                65 if cell.value in {"Title", "Author", "Journal"} else 24
            )
    workbook.save(path)
    unique_count = len({row[-2] for row in rows})
    print(f"Exported {unique_count} unique accessions, {len(rows)} links, "
          f"and {len(by_reference)} references to {path}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--subtype-profile-workbook", required=True, type=Path)
    parser.add_argument("--profile-input-workbook", required=True, type=Path)
    parser.add_argument("--output-csv", required=True, type=Path)
    parser.add_argument("--reference-workbook", required=True, type=Path)
    parser.add_argument("--reference-sheet", required=True)
    args = parser.parse_args()
    count = export(args.subtype_profile_workbook, args.profile_input_workbook, args.output_csv,
                   args.reference_workbook, args.reference_sheet)
    print(f"Exported {count} accession/RefID links to {args.output_csv}")


if __name__ == "__main__":
    main()
