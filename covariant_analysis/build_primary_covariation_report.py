#!/usr/bin/env python3
"""Create one presentation-ready primary RAS--RAS covariation report CSV.

The script combines the all-sequences RAS--RAS scan for every gene, retains a
position pair when it is significant (q <= alpha) under either weighting, and
places the two weighting summaries on one row. It does not replace the full
analysis result files, which retain every tested pair and diagnostic columns.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
from pathlib import Path


INPUT_SUFFIX = "_covariation_all_sequences_ras_ras.csv"

# The CSV retains machine-friendly field names; these labels are used in the workbook.
EXCEL_FIELD_LABELS = {
    "position_pair": "Position pair",
    "pos1": "Position 1",
    "pos2": "Position 2",
    "n_subtypes": "Eligible subtypes",
    "n_sequences": "Usable sequences",
    "significant_weighting": "Significant analysis",
    "sequence_weighted_adjusted_mi_bits": "Biased by subtype sequence count: adjusted MI (bits)",
    "sequence_weighted_p_value": "Biased by subtype sequence count: P value",
    "sequence_weighted_q_value": "Biased by subtype sequence count: q value",
    "subtype_balanced_adjusted_mi_bits": "Equal-subtype-weighted adjusted MI (bits)",
    "subtype_balanced_p_value": "Equal-subtype-weighted P value",
    "subtype_balanced_q_value": "Equal-subtype-weighted q value",
}

EXCEL_ANALYSIS_LABELS = {
    "sequence_weighted": "Biased by subtype sequence count",
    "subtype_balanced": "Equal-subtype weighted",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "inputs", nargs="*", type=Path,
        help=f"Primary result CSVs. Defaults to every *{INPUT_SUFFIX} in this directory.",
    )
    parser.add_argument("--alpha", type=float, default=0.05,
                        help="Include pairs significant at q <= this value (default: 0.05).")
    parser.add_argument("--include-nonsignificant", action="store_true",
                        help="Include all tested pairs, rather than only pairs significant under either weighting.")
    parser.add_argument("--output", type=Path,
                        help="Output CSV path (default: covariation_primary_ras_ras_report.csv beside this script).")
    parser.add_argument("--excel-output", type=Path,
                        help="Output XLSX path (default: covariation_primary_ras_ras_report.xlsx beside this script).")
    return parser.parse_args()


def gene_for(path: Path) -> str:
    match = re.search(r"(?:^|_)(NS3|NS5A|NS5B)(?:_|$)", path.name.upper())
    if not match:
        raise ValueError(f"Cannot identify gene from {path.name}")
    return match.group(1)


def parse_number(value: str | None) -> float | None:
    if value is None or not value.strip():
        return None
    return float(value)


def format_report_number(value: float | None) -> str:
    """Keep two significant digits for values below one, without scientific notation."""
    if value is None or math.isnan(value):
        return ""
    if value == 0:
        return "0"
    magnitude = abs(value)
    if magnitude < 1:
        # For 0.01234 this produces 0.012; for 0.5, 0.50; for 0.001, 0.0010.
        decimal_places = max(2, 1 - math.floor(math.log10(magnitude)))
        return f"{value:.{decimal_places}f}"
    return f"{value:.2f}"


def read_results(path: Path, alpha: float, include_nonsignificant: bool) -> list[dict[str, object]]:
    required = {
        "pos1", "pos2", "weighting", "n_subtypes", "n_sequences",
        "observed_mi_bits", "adjusted_mi_bits", "p_value", "q_value",
    }
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        fields = set(reader.fieldnames or [])
        if missing := required - fields:
            raise ValueError(f"{path.name} is missing columns: {', '.join(sorted(missing))}")
        rows = list(reader)

    grouped: dict[tuple[int, int], dict[str, dict[str, str]]] = {}
    for row in rows:
        key = (int(row["pos1"]), int(row["pos2"]))
        grouped.setdefault(key, {})[row["weighting"]] = row

    output: list[dict[str, object]] = []
    gene = gene_for(path)
    for (pos1, pos2), summaries in grouped.items():
        sequence = summaries.get("sequence_weighted", {})
        balanced = summaries.get("subtype_balanced", {})
        sequence_q = parse_number(sequence.get("q_value"))
        balanced_q = parse_number(balanced.get("q_value"))
        significant = [
            label for label, q_value in (("sequence_weighted", sequence_q),
                                        ("subtype_balanced", balanced_q))
            if q_value is not None and q_value <= alpha
        ]
        if not include_nonsignificant and not significant:
            continue
        source = sequence or balanced
        output.append({
            "gene": gene,
            "position_pair": f"{pos1}-{pos2}",
            "pos1": pos1,
            "pos2": pos2,
            "n_subtypes": source.get("n_subtypes", ""),
            "n_sequences": source.get("n_sequences", ""),
            "significant_weighting": "; ".join(significant),
            "sequence_weighted_adjusted_mi_bits": format_report_number(
                parse_number(sequence.get("adjusted_mi_bits"))),
            "sequence_weighted_p_value": format_report_number(parse_number(sequence.get("p_value"))),
            "sequence_weighted_q_value": format_report_number(sequence_q),
            "subtype_balanced_adjusted_mi_bits": format_report_number(
                parse_number(balanced.get("adjusted_mi_bits"))),
            "subtype_balanced_p_value": format_report_number(parse_number(balanced.get("p_value"))),
            "subtype_balanced_q_value": format_report_number(balanced_q),
        })
    return output


def smallest_q(row: dict[str, object]) -> float:
    values = [
        parse_number(str(row["sequence_weighted_q_value"])),
        parse_number(str(row["subtype_balanced_q_value"])),
    ]
    return min(value for value in values if value is not None) if any(values) else float("inf")


def exclude_low_sequence_outliers(report_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    """Exclude a gene's smallest row when it has under half the next-smallest count."""
    by_gene: dict[str, list[dict[str, object]]] = {}
    for row in report_rows:
        by_gene.setdefault(str(row["gene"]), []).append(row)

    retained: list[dict[str, object]] = []
    for gene_rows in by_gene.values():
        ranked = sorted(gene_rows, key=lambda row: int(str(row["n_sequences"])))
        if len(ranked) >= 2:
            smallest = int(str(ranked[0]["n_sequences"]))
            second_smallest = int(str(ranked[1]["n_sequences"]))
            if smallest < second_smallest / 2:
                ranked = ranked[1:]
        retained.extend(ranked)
    return retained


def excel_number_format(value: str) -> str:
    """Return a fixed-decimal format matching the presentation CSV value."""
    if not value or "." not in value:
        return "0"
    return "0." + "0" * len(value.split(".", maxsplit=1)[1])


def excel_value(field: str, value: object) -> object:
    """Use reader-friendly analysis names in the workbook only."""
    if field != "significant_weighting" or not value:
        return value
    return "; ".join(
        EXCEL_ANALYSIS_LABELS.get(label, label) for label in str(value).split("; ")
    )


def write_excel_report(path: Path, report_rows: list[dict[str, object]], fields: list[str]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as error:
        raise SystemExit(
            "openpyxl is required for the Excel workbook. Use the project environment: "
            "covariant_analysis/.venv/bin/python build_primary_covariation_report.py"
        ) from error

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet_fields = [field for field in fields if field != "gene"]
    numeric_fields = {
        "pos1", "pos2", "n_subtypes", "n_sequences",
        "sequence_weighted_adjusted_mi_bits", "sequence_weighted_p_value", "sequence_weighted_q_value",
        "subtype_balanced_adjusted_mi_bits", "subtype_balanced_p_value", "subtype_balanced_q_value",
    }
    header_font = Font(bold=True)
    by_gene: dict[str, list[dict[str, object]]] = {gene: [] for gene in ("NS3", "NS5A", "NS5B")}
    for row in report_rows:
        by_gene.setdefault(str(row["gene"]), []).append(row)

    read_me = workbook.create_sheet(title="Read Me")
    read_me.append(["Covariation primary RAS--RAS report"])
    read_me["A1"].font = Font(bold=True, size=14)
    read_me.append([])
    read_me.append(["How to use this workbook"])
    read_me["A3"].font = Font(bold=True)
    read_me.append([
        "Each gene sheet contains RAS--RAS position pairs significant at the selected q-value cutoff "
        "under sequence-weighted analysis, subtype-balanced analysis, or both."
    ])
    read_me.append([
        "Sequence-weighted gives a subtype more influence when it has more usable sequences. "
        "For example, subtype 1a contributes more than a subtype with fewer sequences."
    ])
    read_me.append([
        "Subtype-balanced gives every adequately sampled subtype equal influence, regardless of its sequence count."
    ])
    read_me.append([
        "Compare both summaries: agreement supports a result across subtype weighting choices; disagreement can indicate "
        "that a large subtype is driving the sequence-weighted result."
    ])
    read_me.append([
        "Quality filter: within each gene sheet, the row with the fewest usable sequences is omitted when it has "
        "fewer than half as many sequences as the second-smallest row."
    ])
    read_me.append([])
    read_me.append(["Column", "Meaning"])
    read_me["A9"].font = Font(bold=True)
    read_me["B9"].font = Font(bold=True)
    definitions = [
        ("position_pair", "The two amino-acid positions tested, written as position 1-position 2."),
        ("pos1; pos2", "The same two amino-acid positions in separate numeric columns."),
        ("n_subtypes", "Number of subtypes with sufficient paired observations and amino-acid variation for this pair."),
        ("n_sequences", "Total usable sequences across the eligible subtypes for this pair."),
        ("Significant analysis", "Shows which analysis had q value at or below the chosen cutoff (usually 0.05). Biased by subtype sequence count means larger subtypes have more influence. Equal-subtype weighted means every eligible subtype has equal influence. Seeing both methods is stronger support that the result is not due only to the weighting choice."),
        ("Biased by subtype sequence count: adjusted MI (bits)", "Association strength after subtracting average shuffled MI. Larger subtypes have more influence."),
        ("Biased by subtype sequence count: P/q value", "Permutation P value and multiple-testing-adjusted q value for the analysis biased by subtype sequence count."),
        ("Equal-subtype-weighted adjusted MI (bits)", "Association strength after subtracting average shuffled MI, with every eligible subtype given equal influence."),
        ("Equal-subtype-weighted P/q value", "Permutation P value and multiple-testing-adjusted q value for the equal-subtype-weighted MI."),
    ]
    for definition in definitions:
        read_me.append(definition)
    read_me.column_dimensions["A"].width = 42
    read_me.column_dimensions["B"].width = 110
    for row in read_me.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_index in range(4, read_me.max_row + 1):
        read_me.row_dimensions[row_index].height = 30
    read_me.freeze_panes = "A9"

    for gene, rows in by_gene.items():
        worksheet = workbook.create_sheet(title=gene)
        worksheet.append([EXCEL_FIELD_LABELS.get(field, field) for field in sheet_fields])
        for cell in worksheet[1]:
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            worksheet.append([excel_value(field, row[field]) for field in sheet_fields])
        for row_index in range(2, worksheet.max_row + 1):
            for column_index, field in enumerate(sheet_fields, start=1):
                cell = worksheet.cell(row_index, column_index)
                if field in numeric_fields and cell.value not in ("", None):
                    if field in {"pos1", "pos2", "n_subtypes", "n_sequences"}:
                        cell.value = int(cell.value)
                        cell.number_format = "#,##0"
                    else:
                        rendered = str(cell.value)
                        cell.value = float(rendered)
                        cell.number_format = excel_number_format(rendered)
                cell.alignment = Alignment(vertical="top")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.row_dimensions[1].height = 32
        for column in worksheet.columns:
            width = min(40, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            worksheet.column_dimensions[column[0].column_letter].width = width

    workbook.move_sheet(read_me, offset=len(workbook.worksheets) - 1)
    workbook.save(path)


def main() -> None:
    args = parse_args()
    script_dir = Path(__file__).resolve().parent
    inputs = args.inputs or sorted(script_dir.glob(f"*{INPUT_SUFFIX}"))
    if not inputs:
        raise SystemExit(
            "No primary RAS--RAS result files found. Run analyze_covariation.py first."
        )
    output_path = args.output or script_dir / "covariation_primary_ras_ras_report.csv"
    excel_output_path = args.excel_output or script_dir / "covariation_primary_ras_ras_report.xlsx"
    report_rows: list[dict[str, object]] = []
    for path in inputs:
        report_rows.extend(read_results(path, args.alpha, args.include_nonsignificant))
    report_rows = exclude_low_sequence_outliers(report_rows)
    report_rows.sort(key=lambda row: (smallest_q(row), row["gene"], int(row["pos1"]), int(row["pos2"])))
    fields = [
        "gene", "position_pair", "pos1", "pos2", "n_subtypes", "n_sequences", "significant_weighting",
        "sequence_weighted_adjusted_mi_bits", "sequence_weighted_p_value", "sequence_weighted_q_value",
        "subtype_balanced_adjusted_mi_bits", "subtype_balanced_p_value", "subtype_balanced_q_value",
    ]
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(report_rows)
    write_excel_report(excel_output_path, report_rows, fields)
    print(f"Wrote {len(report_rows)} presentation rows to {output_path}")
    print(f"Wrote Excel workbook with NS3, NS5A, and NS5B sheets to {excel_output_path}")


if __name__ == "__main__":
    main()
